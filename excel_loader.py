"""Generic reader for the "Dr. Vijay Malik Screener Excel Template".

Reads the raw `Data Sheet` tab (screener.in-exported P&L / balance sheet /
cash flow / price, one column per fiscal year) and computes the ratios the
14-module stock research framework needs directly from those raw figures.

The template's own `Dr Vijay Malik Analysis` tab holds the same ratios as
formulas, but a workbook that hasn't been recalculated and saved inside
Excel has no cached results for them (openpyxl with data_only=True then
returns None) -- so we recompute from `Data Sheet` ourselves. This is also
more portable: it doesn't depend on Excel's calculation state at all, only
on the raw figures screener.in always populates.

This module is company-agnostic: point it at any export that uses the same
template shape (company name in `Data Sheet!B1`, same row layout) and it
works unmodified.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl

# Row anchors in the "Data Sheet" tab (fixed by the template, same for every company).
_PL_HEADER_ROW = 16
_PL_ROWS = {
    "sales": 17,
    "raw_material": 18,
    "change_in_inventory": 19,
    "power_fuel": 20,
    "other_mfr_exp": 21,
    "employee_cost": 22,
    "selling_admin": 23,
    "other_expenses": 24,
    "other_income": 25,
    "depreciation": 26,
    "interest": 27,
    "pbt": 28,
    "tax": 29,
    "net_profit": 30,
    "dividend": 31,
}
_BS_HEADER_ROW = 56
_BS_ROWS = {
    "equity_share_capital": 57,
    "reserves": 58,
    "borrowings": 59,
    "net_block": 62,
    "cwip": 63,
    "investments": 64,
    "receivables": 67,
    "inventory": 68,
    "cash_and_bank": 69,
    "num_equity_shares": 70,
    "new_bonus_shares": 71,
    "face_value_history": 72,
}
_CF_HEADER_ROW = 81
_CF_ROWS = {
    "cfo": 82,
    "cfi": 83,
    "cff": 84,
}
_MAX_YEAR_COL = 11  # columns B..K = up to 10 annual periods


def _year_columns(ws, header_row: int) -> list[int]:
    """Column indices (1-based) that hold an actual period date in this block's header row."""
    cols = []
    for col in range(2, _MAX_YEAR_COL + 1):
        if isinstance(ws.cell(row=header_row, column=col).value, datetime):
            cols.append(col)
    return cols


def _series(ws, row: int, cols: list[int]) -> list[float | None]:
    return [ws.cell(row=row, column=c).value for c in cols]


@dataclass
class FinancialData:
    company_name: str
    fiscal_year_ends: list[datetime]
    sales: list[float]
    net_profit: list[float]
    dividend: list[float]
    pbt: list[float]
    tax: list[float]
    interest: list[float]
    depreciation: list[float]
    other_income: list[float]
    operating_profit: list[float]  # Sales minus operating expense lines, excludes other income
    ebit: list[float]  # PBT + Interest
    cfo: list[float]
    capex_approx: list[float]  # delta(Net Block + CWIP) + Depreciation
    equity: list[float]  # Equity Share Capital + Reserves
    debt: list[float]  # Borrowings, 0 if not disclosed
    receivables: list[float]
    inventory: list[float]
    cash_and_bank: list[float]
    investments: list[float]
    net_block: list[float]
    cwip: list[float]
    num_shares: list[float]
    new_bonus_shares: list[float]  # bonus shares issued that year, 0 if none
    face_value_history: list[float]  # per-year face value; a drop signals a stock split
    current_price: float
    market_cap: float
    face_value: float

    # -- helpers -------------------------------------------------------
    def _avg(self, series: list[float], idx: int = -1) -> float:
        """Average of the value at idx and the one before it (Vijay-Malik-style 'Avg' figures)."""
        a = series[idx]
        b = series[idx - 1] if len(series) > 1 else series[idx]
        return (a + b) / 2

    @property
    def latest_year(self) -> datetime:
        return self.fiscal_year_ends[-1]

    def revenue_cagr(self, years: int = 3) -> float | None:
        if len(self.sales) <= years:
            return None
        start, end = self.sales[-1 - years], self.sales[-1]
        if start <= 0:
            return None
        return (end / start) ** (1 / years) - 1

    def fcf(self, idx: int = -1) -> float:
        return self.cfo[idx] - self.capex_approx[idx]

    def fcf_to_net_income(self, idx: int = -1) -> float | None:
        ni = self.net_profit[idx]
        return self.fcf(idx) / ni if ni else None

    def ebit_to_interest(self, idx: int = -1) -> float | None:
        """None signals debt-free / no meaningful interest expense (report as infinite)."""
        interest = self.interest[idx]
        if not interest:
            return None
        return self.ebit[idx] / interest

    def interest_coverage(self, idx: int = -1) -> float | None:
        interest = self.interest[idx]
        if not interest:
            return None
        return self.operating_profit[idx] / interest

    def roe(self, idx: int = -1) -> float:
        return self.net_profit[idx] / self._avg(self.equity, idx)

    def roce(self, idx: int = -1) -> float:
        capital_employed = [e + d for e, d in zip(self.equity, self.debt)]
        return self.ebit[idx] / self._avg(capital_employed, idx)

    def roic(self, idx: int = -1, tax_rate: float = 0.25) -> float:
        """NOPAT / average invested capital, excluding surplus cash & investments."""
        invested_capital = [
            e + d - c - inv
            for e, d, c, inv in zip(self.equity, self.debt, self.cash_and_bank, self.investments)
        ]
        nopat = self.ebit[idx] * (1 - tax_rate)
        return nopat / self._avg(invested_capital, idx)

    def debt_to_equity(self, idx: int = -1) -> float:
        return self.debt[idx] / self.equity[idx]

    def is_debt_free(self, idx: int = -1) -> bool:
        return self.debt[idx] == 0

    def cfo_to_pat(self, idx: int = -1) -> float | None:
        pat = self.net_profit[idx]
        return self.cfo[idx] / pat if pat else None

    def cfo_to_pat_avg(self, years: int = 3) -> float | None:
        idxs = range(-years, 0)
        ratios = [self.cfo_to_pat(i) for i in idxs if len(self.net_profit) >= abs(i)]
        ratios = [r for r in ratios if r is not None]
        return sum(ratios) / len(ratios) if ratios else None

    def receivable_days(self, idx: int = -1) -> float:
        return self.receivables[idx] / self.sales[idx] * 365

    def dividend_payout(self, idx: int = -1) -> float:
        pat = self.net_profit[idx]
        return self.dividend[idx] / pat if pat else 0.0

    def retained_earnings(self, idx: int = -1) -> float:
        return self.net_profit[idx] - self.dividend[idx]

    def unexplained_share_increase_pct(self, idx: int = -1) -> float:
        """Share-count growth NOT explained by a face-value split or a recorded bonus issue.

        A stock split (face value drops, e.g. Rs 10 -> Rs 2) or a bonus issue can raise the
        raw share count several-fold without diluting existing holders at all; a naive
        "did share count go up" check would misread either as dilution. This nets both out
        so what's left is genuine dilution (fresh equity issuance, ESOP exercises, etc.).
        """
        if len(self.num_shares) < abs(idx) + 1:
            return 0.0
        prior = self.num_shares[idx - 1]
        if not prior:
            return 0.0
        split_ratio = 1.0
        if self.face_value_history[idx] and self.face_value_history[idx - 1]:
            split_ratio = self.face_value_history[idx - 1] / self.face_value_history[idx]
        expected = prior * split_ratio + self.new_bonus_shares[idx]
        if not expected:
            return 0.0
        return (self.num_shares[idx] - expected) / expected

    def consecutive_dividend_years(self) -> int:
        count = 0
        for d in reversed(self.dividend):
            if d and d > 0:
                count += 1
            else:
                break
        return count

    def is_profitable(self, idx: int = -1) -> bool:
        return self.net_profit[idx] > 0

    def revenue_growing(self, idx: int = -1) -> bool:
        if len(self.sales) < abs(idx) + 1:
            return False
        return self.sales[idx] > self.sales[idx - 1]

    def losses_improving(self, idx: int = -1) -> bool:
        """Only meaningful when not profitable; True if the loss narrowed vs the prior year."""
        if len(self.net_profit) < abs(idx) + 1:
            return False
        return self.net_profit[idx] > self.net_profit[idx - 1]

    def eps(self, idx: int = -1) -> float:
        return self.net_profit[idx] * 1e7 / self.num_shares[idx]  # net_profit is in Cr

    def pe_ratio(self, idx: int = -1) -> float | None:
        eps = self.eps(idx)
        return self.current_price / eps if eps else None


def load_financials(xlsx_path: str | Path) -> FinancialData:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ds = wb["Data Sheet"]

    company_name = ds.cell(row=1, column=2).value
    current_price = ds.cell(row=8, column=2).value
    market_cap = ds.cell(row=9, column=2).value
    face_value = ds.cell(row=7, column=2).value

    pl_cols = _year_columns(ds, _PL_HEADER_ROW)
    bs_cols = _year_columns(ds, _BS_HEADER_ROW)
    cf_cols = _year_columns(ds, _CF_HEADER_ROW)

    def pl(name: str) -> list[float]:
        return [v or 0.0 for v in _series(ds, _PL_ROWS[name], pl_cols)]

    def bs(name: str) -> list[float]:
        return [v or 0.0 for v in _series(ds, _BS_ROWS[name], bs_cols)]

    def cf(name: str) -> list[float]:
        return [v or 0.0 for v in _series(ds, _CF_ROWS[name], cf_cols)]

    fiscal_year_ends = [ds.cell(row=_PL_HEADER_ROW, column=c).value for c in pl_cols]

    sales = pl("sales")
    interest = pl("interest")
    depreciation = pl("depreciation")
    other_income = pl("other_income")
    pbt = pl("pbt")

    # Verified against the workbook's own figures: screener.in stores "Change in
    # Inventory" such that it must be SUBTRACTED (not added) to reconstruct cost of
    # goods sold -- confirmed by reproducing exact historical PBT figures.
    opex = [
        rm - chg_inv + pf + omfr + emp + sa + oe
        for rm, chg_inv, pf, omfr, emp, sa, oe in zip(
            pl("raw_material"),
            pl("change_in_inventory"),
            pl("power_fuel"),
            pl("other_mfr_exp"),
            pl("employee_cost"),
            pl("selling_admin"),
            pl("other_expenses"),
        )
    ]
    operating_profit = [s - x for s, x in zip(sales, opex)]
    ebit = [p + i for p, i in zip(pbt, interest)]  # EBIT = PBT + Interest

    net_block = bs("net_block")
    cwip = bs("cwip")
    net_fixed = [nb + cw for nb, cw in zip(net_block, cwip)]
    capex_approx = [
        (net_fixed[i] - net_fixed[i - 1] if i > 0 else net_fixed[i]) + depreciation[i]
        for i in range(len(net_fixed))
    ]

    equity = [sc + r for sc, r in zip(bs("equity_share_capital"), bs("reserves"))]

    return FinancialData(
        company_name=company_name,
        fiscal_year_ends=fiscal_year_ends,
        sales=sales,
        net_profit=pl("net_profit"),
        dividend=pl("dividend"),
        pbt=pbt,
        tax=pl("tax"),
        interest=interest,
        depreciation=depreciation,
        other_income=other_income,
        operating_profit=operating_profit,
        ebit=ebit,
        cfo=cf("cfo"),
        capex_approx=capex_approx,
        equity=equity,
        debt=bs("borrowings"),
        receivables=bs("receivables"),
        inventory=bs("inventory"),
        cash_and_bank=bs("cash_and_bank"),
        investments=bs("investments"),
        net_block=net_block,
        cwip=cwip,
        num_shares=bs("num_equity_shares"),
        new_bonus_shares=bs("new_bonus_shares"),
        face_value_history=bs("face_value_history"),
        current_price=current_price,
        market_cap=market_cap,
        face_value=face_value,
    )
